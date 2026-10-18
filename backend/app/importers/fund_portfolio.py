"""A fund's portfolio holdings from the AMC's monthly portfolio / factsheet file (uploaded by you).

AMCs publish one Excel file a month, often with a sheet per scheme. Pick the sheet for your fund
(or save it as its own file); columns are matched fuzzily ("Name of the Instrument", "ISIN",
"% to Net Assets"), with the manual column-mapping step as a fallback. Used for fund overlap.
"""

import io
from dataclasses import dataclass, field
from datetime import date

from openpyxl import load_workbook
from sqlmodel import Session, col, delete, select

from app.importers import groww
from app.importers.groww import NeedsMapping, UserMapping
from app.importers.reference import scheme_resolver
from app.importers.tabular import (
    Grid,
    ImportFileError,
    cell_text,
    parse_number,
    read_grid,
    sheet_rows,
)
from app.models.tables import FundPortfolioRow, FundPortfolioUpload

SKIP = ("total", "sub total", "subtotal", "grand total", "net assets", "net current assets")


@dataclass
class FundPortfolioImported:
    upload_id: int
    amfi_code: str
    scheme_name: str
    rows: int
    total_weight_pct: float
    warnings: list[str] = field(default_factory=list)


def _sheet(content: bytes, sheet: str) -> Grid:
    try:
        wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    except Exception as e:
        raise ImportFileError(f"Could not open the Excel file: {e}") from e
    try:
        names = {ws.title.strip().lower(): ws for ws in wb.worksheets}
        ws = names.get(sheet.strip().lower())
        if ws is None:
            raise ImportFileError(
                f"No sheet named '{sheet}'. Sheets in this file: "
                + ", ".join(w.title for w in wb.worksheets)
            )
        return sheet_rows(ws)
    finally:
        wb.close()


def import_fund_portfolio(
    session: Session,
    content: bytes,
    filename: str,
    amfi_code: str,
    *,
    as_of: date | None = None,
    sheet: str | None = None,
    user: UserMapping | None = None,
) -> FundPortfolioImported | NeedsMapping:
    scheme = scheme_resolver().by_code.get(amfi_code.strip())
    if scheme is None:
        raise ImportFileError(f"AMFI scheme code '{amfi_code}' not found")
    if sheet and filename.lower().endswith((".xlsx", ".xlsm")):
        grid = _sheet(content, sheet)
    else:
        grid = read_grid(content, filename)
    mapping = groww.resolve_mapping(session, grid, "fund_portfolio", user)
    if isinstance(mapping, NeedsMapping):
        return mapping
    rows: list[tuple[str, str | None, float]] = []
    for i in range(mapping.header_row + 1, len(grid)):
        r = grid[i]
        rec = {k: (r[c] if c < len(r) else None) for k, c in mapping.mapping.items()}
        name = cell_text(rec.get("name"))
        w = parse_number(rec.get("weight_pct"))
        if not name or w is None or w <= 0:
            continue
        if name.lower().strip(" :-").startswith(SKIP):
            continue
        isin = cell_text(rec.get("isin")).upper() or None
        rows.append((name, isin if isin and len(isin) == 12 else None, w))
    if not rows:
        raise ImportFileError("No holdings with a weight were found in this file.")
    warnings: list[str] = []
    total = sum(w for _, _, w in rows)
    if total <= 1.5:  # weights given as fractions (0.0523) rather than percent
        rows = [(n, i, w * 100) for n, i, w in rows]
        total *= 100
    if total > 110:
        warnings.append(
            f"Weights add up to {total:.0f}%: the file may include several schemes or subtotals. "
            "Pick the sheet for this fund only."
        )
    # Replace this fund's previous upload.
    old = session.exec(
        select(FundPortfolioUpload.id).where(FundPortfolioUpload.amfi_code == scheme.code)
    ).all()
    if old:
        session.exec(delete(FundPortfolioRow).where(col(FundPortfolioRow.upload_id).in_(old)))
        session.exec(delete(FundPortfolioUpload).where(col(FundPortfolioUpload.id).in_(old)))
    up = FundPortfolioUpload(
        amfi_code=scheme.code,
        filename=filename,
        as_of=as_of,
        rows=len(rows),
        total_weight_pct=round(total, 2),
    )
    session.add(up)
    session.flush()
    for nm, isin_, w in rows:
        session.add(
            FundPortfolioRow(upload_id=up.id or 0, name=nm, isin=isin_, weight_pct=round(w, 4))
        )
    session.commit()
    if as_of is None:
        warnings.append("No 'as of' date given: overlap will be marked approximate.")
    return FundPortfolioImported(
        up.id or 0, scheme.code, scheme.name, len(rows), round(total, 2), warnings
    )
