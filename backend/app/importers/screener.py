"""Screener.in uploads (never scraped: you export the file yourself and upload it).

Two kinds of export are understood:

* **Screen results** (CSV or Excel): one row per company, with whatever columns your screen
  shows. Ownership data (promoter holding, pledge, FII/DII and their changes) comes from here.
  Columns are matched with the same fuzzy column mapper as the Groww importer.
* **Company export** ("Export to Excel" on a company page): its "Data Sheet" tab holds about
  10 years of profit & loss, balance sheet and cash flow plus recent quarters, in ₹ crore.

Rows are matched to NSE symbols by NSE code when present, else by company name.
"""

import io
import json
import logging
from dataclasses import dataclass, field
from datetime import date, datetime
from typing import Any

from openpyxl import load_workbook
from sqlmodel import Session

from app.importers import groww
from app.importers.groww import NeedsMapping, UserMapping
from app.importers.reference import ticker_overrides, ticker_resolver
from app.importers.tabular import (
    Cell,
    Grid,
    ImportFileError,
    cell_text,
    parse_date,
    parse_number,
    sheet_rows,
)
from app.models.tables import ScreenerData, ScreenerUpload, utcnow

logger = logging.getLogger(__name__)

DATA_SHEET = "data sheet"


@dataclass
class ScreenerImported:
    kind: str  # "screen" | "company"
    upload_id: int
    matched: list[tuple[str, str]]  # (symbol, name as in the file)
    unmatched: list[str]  # names that couldn't be matched to an NSE symbol
    warnings: list[str] = field(default_factory=list)


# ---------------------------------------------------------------------------- company Data Sheet

# Data Sheet row label (lower-case) → canonical field, per section.
PL_ROWS = {
    "sales": "revenue",
    "net profit": "net_income",
    "interest": "interest_expense",
    "depreciation": "depreciation",
    "profit before tax": "pretax_income",
    "dividend amount": "dividends_paid",
}
QUARTER_ROWS = {"sales": "revenue", "net profit": "net_income", "operating profit": "ebitda"}
BS_ROWS = {
    "borrowings": "total_debt",
    "cash & bank": "cash",
    "no. of equity shares": "shares",
}
CF_ROWS = {"cash from operating activity": "operating_cash_flow"}
SECTIONS = {
    "profit & loss": "pl",
    "quarters": "quarters",
    "balance sheet": "bs",
    "cash flow": "cf",
    "price": None,
    "derived": None,
    "meta": "meta",
}


def _label(v: Cell) -> str:
    return " ".join(cell_text(v).lower().rstrip(":").split())


def _date(v: Cell) -> date | None:
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    return parse_date(v)


def parse_data_sheet(grid: Grid) -> tuple[str | None, dict[str, Any]]:
    """(company name, {"annual": {iso: {field: ₹ Cr}}, "quarterly": {...}}) from a Data Sheet."""
    name: str | None = None
    section: str | None = None
    dates: list[date | None] = []
    raw: dict[str, dict[str, list[float | None]]] = {"pl": {}, "quarters": {}, "bs": {}, "cf": {}}
    sec_dates: dict[str, list[date | None]] = {}
    totals_seen = 0
    for row in grid:
        if not row:
            continue
        lab = _label(row[0])
        if not lab:
            continue
        if lab == "company name":
            name = cell_text(row[1]) if len(row) > 1 else None
            continue
        sec = next((v for k, v in SECTIONS.items() if lab.startswith(k)), "missing")
        if sec != "missing" and len([c for c in row[1:] if cell_text(c)]) == 0:
            section, dates, totals_seen = sec, [], 0
            continue
        if section is None or section == "meta":
            continue
        if lab == "report date":
            dates = [_date(c) for c in row[1:]]
            sec_dates[section] = dates
            continue
        if section not in raw or not dates:
            continue
        vals = [parse_number(c) for c in row[1 : 1 + len(dates)]]
        key = lab
        if section == "bs" and lab == "total":
            totals_seen += 1
            key = "total" if totals_seen == 1 else "total (assets)"
        raw[section].setdefault(key, vals)

    def by_date(section: str) -> dict[str, dict[str, float]]:
        out: dict[str, dict[str, float]] = {}
        for i, d in enumerate(sec_dates.get(section, [])):
            if d is None:
                continue
            vals: dict[str, float] = {}
            for k, v in raw[section].items():
                x = v[i] if i < len(v) else None
                if x is not None:
                    vals[k] = x
            out[d.isoformat()] = vals
        return out

    pl, bs, cf, qt = by_date("pl"), by_date("bs"), by_date("cf"), by_date("quarters")
    annual: dict[str, dict[str, float]] = {}
    for d in sorted(set(pl) | set(bs) | set(cf)):
        p, b, c = pl.get(d, {}), bs.get(d, {}), cf.get(d, {})
        row_out: dict[str, float] = {}
        for src, mapping in ((p, PL_ROWS), (b, BS_ROWS), (c, CF_ROWS)):
            for lab, fld in mapping.items():
                if lab in src:
                    row_out[fld] = src[lab]
        if "profit before tax" in p and "interest" in p:
            row_out["ebit"] = p["profit before tax"] + p["interest"]
            if "depreciation" in p:
                row_out["ebitda"] = row_out["ebit"] + p["depreciation"]
            if "other income" in p:
                row_out["operating_income"] = row_out["ebit"] - p["other income"]
        if "sales" in p and "raw material cost" in p:
            row_out["gross_profit"] = (
                p["sales"] - p["raw material cost"] + p.get("change in inventory", 0.0)
            )
        if "equity share capital" in b and "reserves" in b:
            row_out["equity"] = b["equity share capital"] + b["reserves"]
        if "total" in b:
            row_out["total_assets"] = b["total"]
        if "borrowings" in b:
            row_out["long_term_debt"] = b["borrowings"]  # Screener doesn't split; used as a proxy
        if row_out:
            annual[d] = row_out
    quarterly = {
        d: {fld: q[lab] for lab, fld in QUARTER_ROWS.items() if lab in q} for d, q in qt.items()
    }
    quarterly = {d: v for d, v in quarterly.items() if v}
    return name, {"annual": annual, "quarterly": quarterly}


def _company_grid(content: bytes) -> Grid | None:
    """The "Data Sheet" tab of a Screener.in company export, or None for other workbooks."""
    try:
        wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    except Exception:
        return None
    try:
        for ws in wb.worksheets:
            if ws.title.strip().lower() == DATA_SHEET:
                return sheet_rows(ws)
        return None
    finally:
        wb.close()


# ---------------------------------------------------------------------------- matching & saving


def _resolve(name: str, nse_code: str | None, session: Session) -> str | None:
    # Screener.in shows short names ("TCS", "ITC") that are often the NSE symbol itself.
    hint = nse_code or (name.strip().upper() if " " not in name.strip() else None)
    res = ticker_resolver().resolve(name, None, hint, ticker_overrides(session))
    return res.symbol if res.symbol and res.exchange == "NSE" else None


def _save(
    session: Session,
    symbol: str,
    upload: ScreenerUpload,
    *,
    fields: dict[str, float] | None,
    statements: dict[str, Any] | None,
) -> None:
    row = session.get(ScreenerData, symbol) or ScreenerData(symbol=symbol, upload_id=upload.id or 0)
    row.upload_id = upload.id or 0
    row.updated_at = utcnow()
    if fields:
        merged = {**json.loads(row.fields_json or "{}"), **fields}
        row.fields_json = json.dumps(merged)
    if statements is not None:
        row.statements_json = json.dumps(statements)
    session.add(row)


def import_screener(
    session: Session,
    content: bytes,
    filename: str,
    user: UserMapping | None = None,
    symbol: str | None = None,
) -> ScreenerImported | NeedsMapping:
    grid = _company_grid(content) if filename.lower().endswith((".xlsx", ".xlsm")) else None
    if grid is not None:
        return _import_company(session, grid, filename, symbol)
    return _import_screen(session, content, filename, user)


def _import_company(
    session: Session, grid: Grid, filename: str, symbol: str | None
) -> ScreenerImported:
    name, statements = parse_data_sheet(grid)
    if not statements["annual"]:
        raise ImportFileError(
            "This workbook has a Data Sheet tab, but no yearly results could be read from it."
        )
    sym = symbol.strip().upper().removesuffix(".NS") if symbol else None
    if sym and sym not in ticker_resolver().by_symbol:
        raise ImportFileError(f"{sym} is not an NSE symbol")
    if sym is None:
        sym = _resolve(name or "", None, session) if name else None
    if sym is None:
        raise ImportFileError(
            f"Could not match '{name or 'this company'}' to an NSE symbol. "
            "Enter the NSE symbol and upload again."
        )
    upload = ScreenerUpload(kind="company", filename=filename, rows=1)
    session.add(upload)
    session.flush()
    _save(session, sym, upload, fields=None, statements=statements)
    session.commit()
    years = len(statements["annual"])
    warnings = []
    if years < 6:
        warnings.append(f"Only {years} years of results in the file; 5-year growth needs 6.")
    return ScreenerImported("company", upload.id or 0, [(sym, name or sym)], [], warnings)


def _import_screen(
    session: Session, content: bytes, filename: str, user: UserMapping | None
) -> ScreenerImported | NeedsMapping:
    from app.importers.tabular import read_grid

    grid = read_grid(content, filename)
    mapping = groww.resolve_mapping(session, grid, "screener", user)
    if isinstance(mapping, NeedsMapping):
        return mapping
    upload = ScreenerUpload(kind="screen", filename=filename)
    session.add(upload)
    session.flush()
    matched: list[tuple[str, str]] = []
    unmatched: list[str] = []
    for i in range(mapping.header_row + 1, len(grid)):
        row = grid[i]
        rec = {k: (row[c] if c < len(row) else None) for k, c in mapping.mapping.items()}
        name = cell_text(rec.get("name"))
        if not name or name.lower().startswith(("median", "total", "average")):
            continue
        nse = cell_text(rec.get("nse_code")) or None
        sym = _resolve(name, nse, session)
        if sym is None:
            unmatched.append(name)
            continue
        fields = {
            k: v
            for k in mapping.mapping
            if k not in ("name", "nse_code", "bse_code")
            and (v := parse_number(rec.get(k))) is not None
        }
        _save(session, sym, upload, fields=fields, statements=None)
        matched.append((sym, name))
    upload.rows = len(matched)
    session.commit()
    warnings = []
    if not any(k in mapping.mapping for k in ("promoter_pct", "pledge_pct", "fii_pct", "dii_pct")):
        warnings.append(
            "No ownership columns found. Add Promoter holding, Pledged percentage and FII/DII "
            "holding (and their changes) to your Screener.in screen to get ownership readings."
        )
    if unmatched:
        warnings.append(
            f"{len(unmatched)} row(s) couldn't be matched to an NSE symbol. Add the 'NSE Code' "
            "column to your screen to make matching exact."
        )
    return ScreenerImported("screen", upload.id or 0, matched, unmatched, warnings)
