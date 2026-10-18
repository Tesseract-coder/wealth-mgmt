"""Read an uploaded XLSX/CSV into a ragged grid of cells, plus value parsers for messy exports."""

import csv
import io
import re
from datetime import date, datetime
from typing import Any

from openpyxl import load_workbook

Cell = str | float | int | date | datetime | None
Grid = list[list[Cell]]


class ImportFileError(ValueError):
    """The uploaded file can't be read; the message is safe to show to the user."""


def read_grid(content: bytes, filename: str) -> Grid:
    """The sheet with the most rows (CSV files have just one)."""
    return max((g for _, g in read_sheets(content, filename)), key=len, default=[])


def read_sheets(content: bytes, filename: str) -> list[tuple[str, Grid]]:
    """Every sheet as (title, grid); a CSV is a single untitled sheet."""
    name = filename.lower()
    if name.endswith((".xlsx", ".xlsm")):
        return _read_xlsx(content)
    if name.endswith((".csv", ".txt")):
        return [("", _read_csv(content))]
    if name.endswith(".xls"):
        raise ImportFileError(
            "Old .xls files are not supported. Open it and 'Save As' .xlsx or .csv, then retry."
        )
    raise ImportFileError(f"Unsupported file type: {filename}. Upload .xlsx or .csv.")


def sheet_rows(ws: Any) -> Grid:
    """All rows of a worksheet. Some exporters (e.g. Zerodha Kite) write a wrong sheet size, which
    read-only mode trusts and would cut the sheet down to one row, so recompute it."""
    if hasattr(ws, "reset_dimensions"):
        ws.reset_dimensions()
    return [list(r) for r in ws.iter_rows(values_only=True)]


def _read_xlsx(content: bytes) -> list[tuple[str, Grid]]:
    try:
        wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    except Exception as e:
        raise ImportFileError(f"Could not open the Excel file: {e}") from e
    try:
        return [(ws.title, sheet_rows(ws)) for ws in wb.worksheets]
    finally:
        wb.close()


def _read_csv(content: bytes) -> Grid:
    for enc in ("utf-8-sig", "cp1252", "latin-1"):
        try:
            text = content.decode(enc)
            break
        except UnicodeDecodeError:
            continue
    try:
        dialect: Any = csv.Sniffer().sniff(text[:4096], delimiters=",;\t")
    except csv.Error:
        dialect = csv.excel
    return [list(r) for r in csv.reader(io.StringIO(text), dialect)]


def cell_text(v: Cell) -> str:
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v).strip()


_NUM_JUNK = re.compile(r"[₹,\s%]|rs\.?|inr", re.I)


def parse_number(v: Cell) -> float | None:
    """'₹1,23,456.78' → 123456.78; '(1,234)' → -1234; '-', '', 'N/A' → None."""
    if v is None or isinstance(v, (date, datetime)):
        return None
    if isinstance(v, (int, float)):
        return float(v)
    s = str(v).strip().replace("−", "-")  # unicode minus
    negative = s.startswith("(") and s.endswith(")")
    s = _NUM_JUNK.sub("", s.strip("()"))
    if s in ("", "-", "--", "na", "n/a", "NA", "N/A", "nil"):
        return None
    try:
        n = float(s)
    except ValueError:
        return None
    return -n if negative else n


_DATE_FORMATS = ("%d-%m-%Y", "%d/%m/%Y", "%Y-%m-%d", "%d-%b-%Y", "%d %b %Y", "%d-%b-%y", "%d/%m/%y")


def parse_date(v: Cell) -> date | None:
    if v is None:
        return None
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    s = str(v).strip()
    for fmt in _DATE_FORMATS:
        try:
            return datetime.strptime(s, fmt).date()
        except ValueError:
            continue
    return None


def parse_isin(v: Cell) -> str | None:
    s = cell_text(v).upper()
    return s if re.fullmatch(r"IN[A-Z0-9]{10}", s) else None
