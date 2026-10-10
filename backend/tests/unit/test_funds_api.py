import json
from datetime import date, timedelta
from pathlib import Path

import httpx
import pytest
from fastapi.testclient import TestClient
from sqlmodel import Session

from app.analysis.mf import holding_cash_flows
from app.data import http
from app.models.tables import MFHolding, MFTransaction, Snapshot
from tests.unit.test_signals_api import _fake_sources

SAMPLES = Path(__file__).resolve().parents[3] / "sample_data"
TODAY = date.today()


def _mfapi(monkeypatch: pytest.MonkeyPatch, rates: dict[str, float] | None = None) -> list[str]:
    """Every scheme: six years of daily NAVs growing at rates[code] % a year (default 12)."""
    asked: list[str] = []

    def handler(req: httpx.Request) -> httpx.Response:
        code = req.url.path.rstrip("/").split("/")[-1]
        asked.append(code)
        rate = (rates or {}).get(code, 12.0)
        days = [TODAY - timedelta(days=i) for i in range(365 * 6)]
        data = [
            {
                "date": d.strftime("%d-%m-%Y"),
                "nav": f"{10 * (1 + rate / 100) ** ((d - days[-1]).days / 365.25):.4f}",
            }
            for d in days
        ]
        body = {
            "meta": {"scheme_code": int(code), "scheme_name": "Fund"},
            "status": "SUCCESS",
            "data": data,
        }
        return httpx.Response(200, content=json.dumps(body))

    http.set_client(httpx.Client(transport=httpx.MockTransport(handler)))
    return asked


def _import(client: TestClient) -> None:
    for url, name in (
        ("/import/groww-stocks", "groww_stocks_holdings.xlsx"),
        ("/import/groww-mf", "groww_mf_holdings.xlsx"),
    ):
        with open(SAMPLES / name, "rb") as f:
            assert client.post(url, files={"file": (name, f)}).status_code == 200


def _codes(client: TestClient) -> dict[str, str]:
    return {
        m["scheme_name"]: m["amfi_code"] for m in client.get("/portfolio/holdings").json()["mf"]
    }


def test_fund_report_regular_plan_switches_to_direct(
    client: TestClient, monkeypatch: pytest.MonkeyPatch
) -> None:
    _import(client)
    _fake_sources(monkeypatch)
    codes = _codes(client)
    regular = codes["Mirae Asset Large Cap Fund Regular Growth"]
    asked = _mfapi(monkeypatch, {regular: 11.0})
    r = client.get(f"/analysis/mf/{regular}")
    assert r.status_code == 200, r.text
    b = r.json()
    assert b["plan"] == "regular" and b["direct_twin"] is not None
    assert b["direct_twin"]["code"] in asked  # the Direct plan's NAVs were fetched too
    assert b["direct_gap_pp"] == pytest.approx(1.0, abs=0.05)  # 12% direct vs 11% regular
    assert b["signal"] == "switch" and "exit load" in b["rule"]
    assert b["benchmark"]["label"] == "Nifty 50"
    rets = {x["years"]: x for x in b["returns"]}
    assert rets[3]["fund"] == pytest.approx(11.0, abs=0.1) and rets[3]["benchmark"] is not None
    assert b["risk"]["volatility"] >= 0 and b["rolling"]
    assert b["holding"]["units"] > 0 and b["holding"]["xirr"] is None
    assert "detailed CAS" in b["holding"]["xirr_basis"]
    assert any("Expense ratio" in u for u in b["unavailable"])
    assert any("price index" in n for n in b["notes"])
    assert "SEBI-registered" in b["disclaimer"]
    assert client.get("/analysis/mf/99999999").status_code == 422
    assert client.get("/analysis/mf/999999").status_code == 404


def test_expense_ratio_input(client: TestClient, monkeypatch: pytest.MonkeyPatch) -> None:
    _import(client)
    _fake_sources(monkeypatch)
    code = _codes(client)["Parag Parikh Flexi Cap Fund Direct Growth"]
    _mfapi(monkeypatch, {code: 60.0})  # comfortably ahead of the fake benchmark
    client.get(f"/analysis/mf/{code}")
    b = client.put(f"/analysis/mf/{code}/expense-ratio", json={"expense_ratio_pct": 1.6}).json()
    assert b["expense_ratio_pct"] == 1.6 and b["signal"] == "review" and "1.60%" in b["rule"]
    b = client.put(f"/analysis/mf/{code}/expense-ratio", json={"expense_ratio_pct": None}).json()
    assert b["expense_ratio_pct"] is None
    assert (
        client.put(f"/analysis/mf/{code}/expense-ratio", json={"expense_ratio_pct": 9}).status_code
        == 422
    )


def _portfolio_csv(rows: list[tuple[str, str, float]]) -> bytes:
    lines = [
        "Monthly portfolio as on 30 Sep 2026",
        "",
        "Name of the Instrument,ISIN,Industry,% to Net Assets",
    ]
    lines += [f"{n},{i},X,{w}" for n, i, w in rows]
    lines += ["Total,,,100", "Net Current Assets,,,1.2"]
    return "\n".join(lines).encode()


def test_fund_portfolio_upload_and_overlap(
    client: TestClient, monkeypatch: pytest.MonkeyPatch
) -> None:
    _import(client)
    _fake_sources(monkeypatch)
    codes = _codes(client)
    a = codes["Parag Parikh Flexi Cap Fund Direct Growth"]
    b = codes["Axis ELSS Tax Saver Fund Direct Growth"]
    _mfapi(monkeypatch, {a: 60.0, b: 60.0})
    common = [("HDFC Bank", "INE040A01034", 30.0), ("ICICI Bank", "INE090A01021", 30.0)]
    for code, extra in ((a, ("ITC", "INE154A01025", 40.0)), (b, ("TCS", "INE467B01029", 40.0))):
        r = client.post(
            "/import/fund-portfolio",
            files={"file": ("p.csv", _portfolio_csv([*common, extra]))},
            data={"amfi_code": code, "as_of": (TODAY - timedelta(days=10)).isoformat()},
        )
        assert r.status_code == 200, r.text
        assert r.json()["rows"] == 3 and r.json()["total_weight_pct"] == 100
    for c in codes.values():
        client.get(f"/analysis/mf/{c}")  # cache NAVs and benchmarks
    body = client.get("/analysis/mf/holdings").json()
    i, j = body["overlap"]["codes"].index(a), body["overlap"]["codes"].index(b)
    assert body["overlap"]["pct"][i][j] == 60.0 and body["overlap"]["pct"][i][i] == 100.0
    assert not body["overlap"]["approximate"]
    others = [k for k, c in enumerate(body["overlap"]["codes"]) if c not in (a, b)]
    assert body["overlap"]["pct"][i][others[0]] is None  # no holdings uploaded for that one
    fa = next(f for f in body["funds"] if f["code"] == a)
    assert fa["signal"] == "review" and "60% of its portfolio overlaps" in fa["rule"]
    assert any("Regular plan" in w for w in body["warnings"])


def test_fund_portfolio_import_edge_cases(client: TestClient) -> None:
    code = "118989"
    frac = "Company,ISIN,Weight\nA Ltd,INE000A01001,0.6\nB Ltd,,0.3\nGrand Total,,0.9\n"
    r = client.post(
        "/import/fund-portfolio", files={"file": ("f.csv", frac.encode())}, data={"amfi_code": code}
    )
    assert r.status_code == 200 and r.json()["total_weight_pct"] == 90
    assert any("as of" in w for w in r.json()["warnings"])
    weird = "Foo,Bar,Baz\n1,2,3\n"
    r = client.post(
        "/import/fund-portfolio",
        files={"file": ("f.csv", weird.encode())},
        data={"amfi_code": code},
    )
    assert (
        r.status_code == 200
        and r.json()["status"] == "needs_mapping"
        and r.json()["kind"] == "fund_portfolio"
    )
    r = client.post(
        "/import/fund-portfolio",
        files={"file": ("f.csv", frac.encode())},
        data={"amfi_code": "999999"},
    )
    assert r.status_code == 400


def test_cash_flows_from_cas_transactions(session: Session) -> None:
    snap = Snapshot(kind="mf", source="cas_cams")
    session.add(snap)
    session.flush()
    h = MFHolding(
        snapshot_id=snap.id or 0,
        scheme_name="X",
        units=10,
        invested=2000,
        first_purchase_date=date(2024, 1, 1),
    )
    session.add(h)
    session.flush()
    for d, t, amt in (
        (date(2024, 1, 1), "PURCHASE", 1000.0),
        (date(2024, 6, 1), "PURCHASE_SIP", 1000.0),
        (date(2024, 6, 1), "STAMP_DUTY_TAX", 0.05),
        (date(2025, 1, 1), "REDEMPTION", -500.0),
        (date(2025, 2, 1), "DIVIDEND_REINVEST", 20.0),
    ):
        session.add(MFTransaction(holding_id=h.id or 0, date=d, type=t, amount=amt))
    session.commit()
    flows, from_txn = holding_cash_flows(session, [h])
    assert from_txn
    assert sorted(flows) == [
        (date(2024, 1, 1), -1000.0),
        (date(2024, 6, 1), -1000.0),
        (date(2024, 6, 1), -0.05),
        (date(2025, 1, 1), 500.0),
    ]
    h2 = MFHolding(
        snapshot_id=snap.id or 0,
        scheme_name="Y",
        units=1,
        invested=900,
        first_purchase_date=date(2023, 5, 1),
    )
    session.add(h2)
    session.commit()
    assert holding_cash_flows(session, [h2]) == ([(date(2023, 5, 1), -900)], False)


def test_portfolio_summary(client: TestClient, monkeypatch: pytest.MonkeyPatch) -> None:
    _import(client)
    _fake_sources(monkeypatch)
    _mfapi(monkeypatch)
    client.post(
        "/import/mf-manual",
        json={"scheme": "118989", "units": 1000, "invested": 9000, "purchase_date": "2023-01-02"},
    )
    for c in _codes(client).values():
        if c:
            client.get(f"/analysis/mf/{c}")
    for t in ("RELIANCE.NS", "ITC.NS", "^NSEI"):
        client.get(f"/analysis/technical/{t}")
    s = client.get("/portfolio/summary").json()
    assert s["total_value"] == pytest.approx(s["stock_value"] + s["fund_value"], abs=0.1)
    assert s["fund_value"] > 0 and s["stock_value"] > 0
    for k in ("by_holding", "by_sector", "by_mcap", "by_asset_class"):
        assert sum(x["weight"] for x in s[k]) == pytest.approx(1.0, abs=0.01), k
    classes = {x["label"] for x in s["by_asset_class"]}
    assert {"Stocks & ETFs", "Equity funds"} <= classes
    assert s["xirr"] is not None and s["xirr"] > 0
    assert 0 < s["xirr_coverage"] < 0.5  # only the manual fund is dated
    assert "Partial" in s["xirr_basis"] and "lump sum" in s["xirr_basis"]
    assert s["beta"] is not None and s["correlation_labels"]
    n = len(s["correlation_labels"])
    assert len(s["correlation"]) == n and all(len(row) == n for row in s["correlation"])
    assert s["performance_dates"] and s["performance_portfolio"][0] == 100.0
    assert "last year is shown" in s["performance_basis"]
    assert len(s["snapshots"]) >= 2
    assert s["day_change"] is not None


def test_portfolio_summary_empty(client: TestClient) -> None:
    s = client.get("/portfolio/summary").json()
    assert s["total_value"] == 0 and s["by_holding"] == [] and s["xirr"] is None
    assert s["invested"] is None and s["day_change"] is None


def test_fund_portfolio_sheet_picker(client: TestClient) -> None:
    import io

    from openpyxl import Workbook

    wb = Workbook()
    wb.active.title = "Index"
    wb.active.append(["Scheme list"])
    ws = wb.create_sheet("PPFCF")
    ws.append(["Name of the Instrument", "ISIN", "% to Net Assets"])
    ws.append(["HDFC Bank Ltd", "INE040A01034", 8.1])
    ws.append(["Sub Total", "", 8.1])
    buf = io.BytesIO()
    wb.save(buf)
    files = {"file": ("amc.xlsx", buf.getvalue())}
    ok = client.post(
        "/import/fund-portfolio", files=files, data={"amfi_code": "118989", "sheet": "ppfcf"}
    )
    assert ok.status_code == 200 and ok.json()["rows"] == 1
    bad = client.post(
        "/import/fund-portfolio", files=files, data={"amfi_code": "118989", "sheet": "nope"}
    )
    assert bad.status_code == 400 and "Sheets in this file: Index, PPFCF" in bad.json()["detail"]
