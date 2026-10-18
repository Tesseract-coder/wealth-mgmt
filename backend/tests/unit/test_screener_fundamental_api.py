import json
from pathlib import Path

import pytest
from fastapi.testclient import TestClient
from sqlmodel import Session

from app.data import yahoo
from app.importers.screener import parse_data_sheet
from app.importers.tabular import read_grid
from app.jobs.runner import get_runner
from app.models.tables import ScreenerData
from tests.conftest import yf_frame
from tests.unit.fundamentals_fixture import fake_fund_download

SAMPLES = Path(__file__).parents[3] / "sample_data"


def upload(client: TestClient, name: str, **data: str) -> dict:  # type: ignore[type-arg]
    with open(SAMPLES / name, "rb") as f:
        r = client.post("/import/screener", files={"file": (name, f)}, data=data)
    assert r.status_code == 200, r.text
    return r.json()


def test_screen_export_matches_and_skips(client: TestClient) -> None:
    body = upload(client, "screener_screen.csv")
    assert body["kind"] == "screen"
    assert {m["symbol"] for m in body["matched"]} == {
        "RELIANCE",
        "TCS",
        "HDFCBANK",
        "ITC",
        "ASIANPAINT",
    }
    assert body["unmatched"] == ["Acme Imaginary Widgets"]  # the "Median" summary row is skipped
    assert any("NSE Code" in w for w in body["warnings"])


def test_screen_export_without_ownership_warns(client: TestClient, tmp_path: Path) -> None:
    p = tmp_path / "plain.csv"
    p.write_text("S.No.,Name,NSE Code,CMP Rs.,P/E,ROCE %\n1,TCS,TCS,3412,24.1,62.5\n")
    with open(p, "rb") as f:
        body = client.post("/import/screener", files={"file": ("plain.csv", f)}).json()
    assert any("No ownership columns" in w for w in body["warnings"])


def test_screen_export_with_unknown_headers_asks_for_mapping(
    client: TestClient, tmp_path: Path
) -> None:
    p = tmp_path / "weird.csv"
    p.write_text("Firm,Thing,Other\nTCS,1,2\n")
    with open(p, "rb") as f:
        body = client.post("/import/screener", files={"file": ("weird.csv", f)}).json()
    assert body["status"] == "needs_mapping" and body["kind"] == "screener"
    assert body["missing_required"] == ["name"]
    with open(p, "rb") as f:
        mapping = json.dumps({"header_row": 0, "columns": {"name": 0}, "remember": False})
        done = client.post(
            "/import/screener", files={"file": ("weird.csv", f)}, data={"mapping": mapping}
        ).json()
    assert done["status"] == "imported" and done["matched"][0]["symbol"] == "TCS"


def test_company_export_parses_ten_years() -> None:
    from openpyxl import load_workbook

    wb = load_workbook(SAMPLES / "screener_company_ITC.xlsx", read_only=True)
    grid = [list(r) for r in wb["Data Sheet"].iter_rows(values_only=True)]
    name, st = parse_data_sheet(grid)
    assert name == "ITC Ltd"
    assert len(st["annual"]) == 10 and len(st["quarterly"]) == 10
    last = st["annual"]["2026-03-31"]
    assert last["equity"] == pytest.approx(1250 + 55000 * 1.07**9, rel=1e-4)
    assert last["ebit"] == pytest.approx(last["pretax_income"] + 40.0)
    assert last["total_assets"] > last["equity"] and "operating_income" in last
    assert read_grid(
        (SAMPLES / "screener_company_ITC.xlsx").read_bytes(), "x.xlsx"
    )  # still readable


def test_company_export_upload_and_symbol_override(
    client: TestClient, session: Session | None = None
) -> None:
    body = upload(client, "screener_company_ITC.xlsx")
    assert body["kind"] == "company" and body["matched"] == [{"symbol": "ITC", "name": "ITC Ltd"}]
    bad = None
    with open(SAMPLES / "screener_company_ITC.xlsx", "rb") as f:
        bad = client.post(
            "/import/screener", files={"file": ("c.xlsx", f)}, data={"symbol": "NOTREAL"}
        )
    assert bad.status_code == 400 and "not an NSE symbol" in bad.json()["detail"]


def _patch_yahoo(monkeypatch: pytest.MonkeyPatch) -> list[str]:
    dl = fake_fund_download()
    monkeypatch.setattr(yahoo, "_fund_download", dl)
    from datetime import date

    monkeypatch.setattr(
        yahoo,
        "_download",
        lambda t, **k: yf_frame(k.get("start") or date(2021, 1, 1), date.today()),
    )
    return dl.calls  # type: ignore[attr-defined,no-any-return]


def test_fundamental_endpoint_yahoo(client: TestClient, monkeypatch: pytest.MonkeyPatch) -> None:
    calls = _patch_yahoo(monkeypatch)
    body = client.get("/analysis/fundamental/LT.NS").json()
    assert body["source"] == "yahoo" and body["label"] == "Fake Industries Ltd"
    assert body["sector"] == "Construction"  # NSE industry for LT from the Nifty 500 list
    metrics = {m["key"]: m for m in body["metrics"]}
    assert metrics["pe"]["basis"].startswith("Market cap ÷ net profit (TTM")
    assert body["fscore"]["score"] == 9
    assert any("Sector medians (0 Construction peer(s)" in u for u in body["unavailable"])
    assert "not investment advice" in body["disclaimer"]
    assert len(calls) == 6
    client.get("/analysis/fundamental/LT.NS")
    assert len(calls) == 6  # cached for fundamentals_stale_after_days
    assert metrics["peg"]["basis"].startswith("Current P/E ÷ expected growth (analyst")
    g = client.get("/analysis/growth-guidance/LT.NS").json()
    assert g["growth_pct"] is None and g["analyst_pct"] == pytest.approx(20.0)
    assert client.put("/analysis/growth-guidance/LT.NS", json={"growth_pct": 40}).json()[
        "growth_pct"
    ] == pytest.approx(40)
    peg = {m["key"]: m for m in client.get("/analysis/fundamental/LT.NS").json()["metrics"]}["peg"]
    assert peg["value"] == pytest.approx(round(metrics["pe"]["value"] / 40, 2), abs=0.01)
    assert "your guidance" in peg["basis"]
    assert (
        client.put("/analysis/growth-guidance/LT.NS", json={"growth_pct": None}).json()[
            "growth_pct"
        ]
        is None
    )


def test_screener_company_statements_take_over(
    client: TestClient, monkeypatch: pytest.MonkeyPatch
) -> None:
    _patch_yahoo(monkeypatch)
    upload(client, "screener_company_ITC.xlsx")
    upload(client, "screener_screen.csv")
    body = client.get("/analysis/fundamental/ITC.NS").json()
    assert body["source"] == "screener"
    assert any("Screener.in upload (10 years)" in n for n in body["notes"])
    assert any("current assets" in n for n in body["notes"])  # filled from Yahoo, said so
    metrics = {m["key"]: m for m in body["metrics"]}
    assert "revenue_cagr_5y" in metrics and "Screener.in" in metrics["revenue_cagr_5y"]["basis"]
    assert metrics["fii_change"]["value"] == 0.4


def test_stale_screener_upload_is_ignored(
    client: TestClient, monkeypatch: pytest.MonkeyPatch
) -> None:
    from datetime import timedelta

    from app.core.db import get_engine
    from app.models.tables import utcnow

    _patch_yahoo(monkeypatch)
    upload(client, "screener_screen.csv")
    with Session(get_engine()) as s:
        row = s.get(ScreenerData, "TCS")
        assert row
        row.updated_at = utcnow() - timedelta(days=200)
        s.add(row)
        s.commit()
    body = client.get("/analysis/fundamental/TCS.NS").json()
    assert any("200 days old" in n for n in body["notes"])
    assert not any(m["key"] == "promoter_pct" for m in body["metrics"])


def test_index_has_no_fundamentals(client: TestClient) -> None:
    body = client.get("/analysis/fundamental/^NSEI").json()
    assert body["metrics"] == [] and body["unavailable"] == [
        "Fundamentals (not applicable to an index)"
    ]


def test_fundamentals_job_with_peers_and_holdings_summary(
    client: TestClient, monkeypatch: pytest.MonkeyPatch
) -> None:
    calls = _patch_yahoo(monkeypatch)
    from app.core import profiles

    p = profiles.get_profiles()
    profiles._profiles = p.model_copy(
        update={"fundamental": p.fundamental.model_copy(update={"sector_max_peers": 6})}
    )
    with open(SAMPLES / "groww_stocks_holdings.xlsx", "rb") as f:
        client.post("/import/groww-stocks", files={"file": f})
    job = client.post("/market/refresh", json={"scope": "fundamentals"}).json()
    assert job["kind"] == "fundamentals_refresh"
    get_runner().wait(job["id"], timeout=60)
    done = client.get(f"/jobs/{job['id']}").json()
    assert done["status"] == "done" and done["total"] > 11  # holdings + peers
    tickers = {c.split(":")[0] for c in calls}
    assert "RELIANCE.NS" in tickers and len(tickers) == done["total"]
    rows = client.get("/analysis/fundamental/holdings").json()
    assert len(rows) == 11
    lt = next(r for r in rows if r["ticker"] == "LT.NS")
    assert lt["pe"] is not None and lt["fscore"] == 9
    status = client.get("/market/status").json()
    assert status["last_fundamentals_job"]["id"] == job["id"]
    assert status["schedule"]["fundamentals"].startswith("Weekly, Sat")
    # Every stock here is the same fake company, so peers give medians: P/E is "in line".
    full = client.get("/analysis/fundamental/LT.NS").json()
    pe = next(m for m in full["metrics"] if m["key"] == "pe")
    assert pe["peers"] >= 5 and pe["sector_median"] is not None


def test_etf_has_no_company_fundamentals(client: TestClient) -> None:
    body = client.get("/analysis/fundamental/NIFTYBEES.NS").json()
    assert body["metrics"] == [] and "ETF" in body["unavailable"][0]
    assert body["freshness"]["status"] == "unavailable"
